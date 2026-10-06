"""``GET /api/file-sheet``: the bounded spreadsheet grid preview."""

from __future__ import annotations

import asyncio
import datetime as _dt
import io
import zipfile
from typing import TYPE_CHECKING, BinaryIO

from aiohttp import web

if TYPE_CHECKING:
    from kiro_crew.dashboard.handlers.files import (
        _MAX_UPLOAD_BYTES,
        _SHEET_MAX_CDIR_ENTRY_BYTES,
        _SHEET_MAX_CELL_CHARS,
        _SHEET_MAX_COLS,
        _SHEET_MAX_EXPANDED_BYTES,
        _SHEET_MAX_MEMBERS,
        _SHEET_MAX_ROWS,
        _SHEET_MAX_SHEETS,
        _SHEET_MAX_TEXT_CHARS,
        ZipInventoryRejected,
        _open_checked_file,
        _OpenDenied,
        _PathProbeBusy,
        _probe_busy_response,
        _run_path_probe,
        _sel,
        logger,
        redact,
        require_owner_dashboard_request,
        vet_zip_inventory_bytes,
    )


class _SheetRefusal(Exception):
    """Deliberate refusal carrying its HTTP status and machine-readable code;
    raised on the worker thread and mapped to a response by api_file_sheet."""

    def __init__(self, status: int, message: str, code: str):
        super().__init__(message)
        self.status = status
        self.code = code


def _sheet_cell_json(value: object) -> object:
    """Serialize one workbook cell value into a JSON-safe primitive."""
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, str):
        # Workbook text is file content leaving the host through the dashboard
        # — same egress class as api_file_read, so the same redaction applies.
        # Redact BEFORE truncating so the scan always sees the complete text,
        # then cap the cell so one shared string cannot bloat every row.
        value = redact(value)
        if len(value) > _SHEET_MAX_CELL_CHARS:
            return value[:_SHEET_MAX_CELL_CHARS] + "…"
        return value
    if isinstance(value, float):
        # NaN/Infinity are rejected by JSON.parse in the browser; the stdlib
        # encoder would happily emit the JS-only tokens.
        if value != value or value in (float("inf"), float("-inf")):
            return str(value)
        return value
    if isinstance(value, _dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (_dt.date, _dt.time)):
        return value.isoformat()
    return redact(str(value))


def _sheet_formula_text(value: object) -> str | None:
    """Return the formula source ("=…") for a formula-pass cell value, else None."""
    text: object = value
    if not (isinstance(text, str) and text.startswith("=")):
        # Array formulas come back as openpyxl ArrayFormula objects carrying .text.
        text = getattr(value, "text", None)
    if isinstance(text, str) and text.startswith("="):
        text = redact(text)
        if len(text) > _SHEET_MAX_CELL_CHARS:
            return text[:_SHEET_MAX_CELL_CHARS] + "…"
        return text
    return None


def _load_sheet_payload(f: BinaryIO, *, max_bytes: int) -> dict:
    """Read, vet, and parse the workbook into the sheet-grid payload.

    Runs ENTIRELY on a worker thread (via asyncio.to_thread) so filesystem
    latency, the first (heavy) openpyxl import, and parse time never stall the
    gateway event loop. Receives the checked-open file object from
    :func:`_open_checked_file` (the shared open-and-check prefix, which owns
    path validation, the sensitive-path gate, and the symlink-refusing open)
    and takes ownership: the file is closed on every path. The bounded-read
    cap is this endpoint's size policy, passed in as *max_bytes*. openpyxl is
    a soft import: absence surfaces as ImportError from this thread and the
    handler maps it to 501.
    """
    with f:
        import openpyxl  # noqa: F401  (probe here, off-loop; parse imports lazily too)

        # Bounded read is the size guard: a pre-check via fstat would race a
        # concurrent writer (the file can grow between the stat and the read,
        # e.g. an agent still generating the workbook), while reading at most
        # cap+1 bytes bounds memory unconditionally.
        data = f.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise _SheetRefusal(413, "file too large", "file_too_large")
    header = data[:4]
    # OOXML spreadsheets are ZIP containers; refuse anything else before
    # openpyxl touches the bytes.
    if not header.startswith(b"PK\x03\x04"):
        raise _SheetRefusal(415, "not an OOXML spreadsheet", "not_a_spreadsheet")
    # Vet the archive's declared inventory before anything inflates it --
    # including ZipFile construction itself, which materializes one ZipInfo
    # per central-directory entry. The EOCD preflight bounds that allocation
    # from the raw bytes; the infolist() pass then bounds what openpyxl can
    # actually expand (zipfile truncates each member at its declared
    # file_size, so the central directory's numbers are authoritative).
    _vet_zip_eocd(data)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        if (
            len(infos) > _SHEET_MAX_MEMBERS
            or sum(i.file_size for i in infos) > _SHEET_MAX_EXPANDED_BYTES
        ):
            raise _SheetRefusal(413, "workbook expands too large", "workbook_expands_too_large")
    return _parse_workbook_grid(data)


def _vet_zip_eocd(data: bytes) -> None:
    """Refuse archives whose end-of-central-directory record declares an
    oversized inventory, BEFORE zipfile.ZipFile is constructed.

    Delegates to the shared vet (kiro_crew.zip_vet) so this endpoint, knowledge
    ingest, and document parsing share one implementation of the preflight --
    only the caps and the error channel stay per-caller. This endpoint's
    observable behaviour is unchanged: a tail with no usable EOCD still reads as
    "not a spreadsheet" (415), an over-cap inventory as an expansion refusal
    (413).
    """
    try:
        vet_zip_inventory_bytes(
            data,
            max_members=_SHEET_MAX_MEMBERS,
            max_cdir_entry_bytes=_SHEET_MAX_CDIR_ENTRY_BYTES,
        )
    except ZipInventoryRejected as exc:
        if exc.reason in ("missing_eocd", "truncated_eocd", "unreadable"):
            raise _SheetRefusal(415, "not an OOXML spreadsheet", "not_a_spreadsheet") from exc
        raise _SheetRefusal(
            413, "workbook expands too large", "workbook_expands_too_large"
        ) from exc


def _parse_workbook_grid(data: bytes) -> dict:
    """Parse xlsx bytes into a JSON-safe sheet grid. Runs on a worker thread.

    The workbook is loaded twice in read-only streaming mode: once with
    data_only=True (formula cells yield the value cached by the writing
    application) and once with data_only=False (formula cells yield the
    formula source). Cells prefer the cached value; when a file carries no
    cache — typical for openpyxl-generated workbooks — the formula text is
    shown instead of an empty cell. Both loads stream the same bytes, so the
    row structures are identical and can be zipped in lockstep.
    """
    import itertools

    from openpyxl import load_workbook

    wb_vals = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    wb_form = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        names = wb_vals.sheetnames
        sheets: list[dict] = []
        # Cumulative post-truncation text budget across the whole workbook:
        # shared strings are stored once but referenced per cell, so archive
        # size caps alone do not bound the JSON response this grid becomes.
        text_chars = 0
        for name in names[:_SHEET_MAX_SHEETS]:
            ws_v, ws_f = wb_vals[name], wb_form[name]
            if not hasattr(ws_v, "iter_rows"):  # chartsheets have no cell grid
                continue
            # Dimension records can lie (some writers emit a stale ref such as
            # A1:A1 for a populated sheet); read-only mode trusts them, so
            # iter_rows would stop early and silently truncate the preview.
            # Force a real scan of each sheet instead.
            if hasattr(ws_v, "reset_dimensions"):
                ws_v.reset_dimensions()
                ws_f.reset_dimensions()
            raw: list[list[object]] = []
            rows_truncated = False
            cols_truncated = False
            paired = zip(ws_v.iter_rows(values_only=True), ws_f.iter_rows(values_only=True))
            for vrow, frow in itertools.islice(paired, _SHEET_MAX_ROWS + 1):
                if len(raw) >= _SHEET_MAX_ROWS:
                    rows_truncated = True
                    # No total is reported: any count derived from workbook
                    # geometry is attacker-influenced (a single sparse row at
                    # index 1e9 makes the read-only reader synthesize a
                    # billion empties), so nothing here iterates past the cap.
                    break
                if len(vrow) > _SHEET_MAX_COLS:
                    cols_truncated = True
                out: list[object] = []
                for vv, fv in list(zip(vrow, frow))[:_SHEET_MAX_COLS]:
                    ftxt = _sheet_formula_text(fv)
                    cell = ftxt if (vv is None and ftxt) else _sheet_cell_json(vv)
                    if isinstance(cell, str):
                        text_chars += len(cell)
                        if text_chars > _SHEET_MAX_TEXT_CHARS:
                            raise _SheetRefusal(
                                413,
                                "workbook text too large to preview",
                                "workbook_text_too_large",
                            )
                    out.append(cell)
                raw.append(out)
            # Trim trailing all-empty rows, then normalize every row to the
            # widest non-empty extent so the client renders a rectangle.
            while raw and all(c is None or c == "" for c in raw[-1]):
                raw.pop()
            width = 0
            for r in raw:
                w = len(r)
                while w and (r[w - 1] is None or r[w - 1] == ""):
                    w -= 1
                width = max(width, w)
            rows = [r[:width] + [None] * (width - len(r[:width])) for r in raw] if width else []
            sheets.append(
                {
                    # Names take the same redact+truncate path as cell text — a
                    # crafted workbook.xml can carry arbitrarily long sheet names.
                    "name": _sheet_cell_json(name),
                    "rows": rows,
                    "truncated_rows": rows_truncated,
                    "truncated_cols": cols_truncated,
                }
            )
        return {
            "sheets": sheets,
            "total_sheets": len(names),
            "truncated_sheets": len(names) > _SHEET_MAX_SHEETS,
        }
    finally:
        wb_vals.close()
        wb_form.close()


async def api_file_sheet(request: web.Request) -> web.Response:
    """GET /api/file-sheet?path=… — parse an OOXML spreadsheet into a JSON cell grid.

    Powers the file viewer's inline xlsx preview. The security prefix is the
    shared :func:`_open_checked_file` (dashboard path validation,
    sensitive-path block, a symlink-refusing open — _open_rb_nofollow: atomic
    O_NOFOLLOW on POSIX, lstat guard on Windows); this endpoint's own policy
    on top is the bounded-read size cap, the zip-expansion caps, and a ZIP
    magic-byte check before openpyxl touches the bytes. All file IO and
    parsing runs on a worker thread so a large workbook cannot stall the
    event loop, and cell text is credential-redacted like every other
    dashboard egress. openpyxl is soft-imported: without it the endpoint
    answers 501 and the frontend degrades to the download card.
    """
    owner_denied = await require_owner_dashboard_request(request, "file_sheet")
    if owner_denied is not None:
        return owner_denied

    def _log(outcome: str, res: str) -> None:
        _sel().log_tool_invocation(
            session_key="dashboard",
            tool_name="file_sheet",
            outcome=outcome,
            resources=res,
        )

    raw_path = request.query.get("path", "")
    # The validated path once the prefix produces one -- exported by the
    # worker callback so the exception handlers log the same SEL resource
    # the success path does.
    res_path = raw_path

    def _open_and_load() -> dict | _OpenDenied:
        """Open-and-check plus parse, in ONE worker-thread hop.

        The checked open file object never crosses back to the event loop:
        every path that opens it also closes it on THIS thread (refusals
        close inside the prefix; the parser's ``with f:`` covers the rest).
        A cancellation of the awaiting task therefore cannot strand an open
        file in a discarded future or finalize one on the loop -- the
        future's result is only ever a payload dict or a typed refusal.
        """
        nonlocal res_path
        checked = _open_checked_file(
            raw_path,
            tool_name="file_sheet",
            log_open_failure=False,
        )
        if isinstance(checked, _OpenDenied):
            return checked
        res_path = checked.path
        return _load_sheet_payload(checked.file, max_bytes=_MAX_UPLOAD_BYTES)

    try:
        result = await _run_path_probe(_open_and_load, transfer=True)
    except asyncio.CancelledError:
        # Shutdown or client disconnect: the access attempt must not vanish
        # from the audit trail. No resource handling here -- the worker
        # callback owns the file's whole lifetime.
        _log("cancelled", res_path)
        raise
    except _PathProbeBusy:
        return _probe_busy_response(resource=res_path, tool_name="file_sheet")
    except ImportError:
        # openpyxl absent: the preview is unavailable, not broken. The probe
        # runs inside the worker thread so even the first heavy import never
        # touches the event loop.
        _log("failure", res_path)
        return web.json_response(
            {"error": "spreadsheet preview unavailable", "code": "preview_unavailable"},
            status=501,
        )
    except _SheetRefusal as refusal:
        # Both refusal kinds map to literal statuses so the response shape
        # stays statically checkable; the carried code names the exact cause.
        _log("denied", res_path)
        if refusal.status == 415:
            return web.json_response({"error": str(refusal), "code": refusal.code}, status=415)
        return web.json_response({"error": str(refusal), "code": refusal.code}, status=413)
    except OSError:
        # Read failure on the already-checked fd. (A symlink never reaches
        # here: the shared prefix refuses it as _OpenDenied("symlink_refused")
        # before the parser sees a file object.)
        _log("failure", res_path)
        return web.json_response({"error": "cannot read file", "code": "read_failed"}, status=500)
    except Exception:
        # openpyxl's failure surface is wide (bad zip members, malformed XML,
        # unexpected workbook parts). Every parse failure degrades to the same
        # client answer, and the frontend falls back to the download card.
        logger.warning("file-sheet: cannot parse workbook %s", res_path, exc_info=True)
        _log("failure", res_path)
        return web.json_response(
            {"error": "cannot parse workbook", "code": "parse_failed"}, status=422
        )
    if isinstance(result, _OpenDenied):
        code, res = result.code, result.path
        if code == "invalid_path":
            _log("denied", res)
            return web.json_response(
                {"error": "invalid or forbidden path", "code": "invalid_path"}, status=400
            )
        if code == "sensitive_path":
            _log("denied", res)
            return web.json_response(
                {"error": "sensitive path blocked", "code": "sensitive_path"}, status=403
            )
        if code == "not_found":
            _log("not_found", res)
            return web.json_response({"error": "not found", "code": "not_found"}, status=404)
        if code == "symlink_refused":
            _log("denied", res)
            return web.json_response(
                {"error": "symlinks not allowed", "code": "symlink_refused"}, status=403
            )
        if code == "file_too_large":
            # Reachable only if this endpoint ever passes fstat_cap; mapped so
            # a policy refusal can never masquerade as the 500 below. (Its
            # size guard today is the bounded read inside _load_sheet_payload.)
            _log("denied", res)
            return web.json_response(
                {"error": "file too large", "code": "file_too_large"}, status=413
            )
        # read_failed: the residual code.
        _log("failure", res)
        return web.json_response({"error": "cannot read file", "code": "read_failed"}, status=500)
    _log("success", res_path)
    return web.json_response(result)
